"""출력물: 엑셀 보고서(4개 시트), 정상·격리 거래 CSV."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Iterable, Mapping, Sequence

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from recon.loader import LoadedFile
from recon.models import FILE_LABELS, FILE_TYPES, ISSUE_COLUMNS
from recon.reconcile import RECON_COLUMNS
from recon.rules.registry import all_rules
from recon.validation import ValidationResult

ISSUE_HEADERS = ("규칙", "심각도", "파일", "원본 행번호", "키", "컬럼", "값", "메시지")

RECON_HEADERS = {
    "account_no": "계좌번호",
    "result": "대사 결과",
    "diff_cause": "차이 원인",
    "summary_row_no": "집계표 행번호",
    "is_target": "집계 대상",
    "in_summary": "집계표 존재",
    "txn_deposit_count": "입금 건수(거래)",
    "sum_deposit_count": "입금 건수(집계표)",
    "deposit_count_diff": "입금 건수 차이",
    "txn_deposit_amount": "입금 금액(거래)",
    "sum_deposit_amount": "입금 금액(집계표)",
    "deposit_amount_diff": "입금 금액 차이",
    "txn_withdrawal_count": "출금 건수(거래)",
    "sum_withdrawal_count": "출금 건수(집계표)",
    "withdrawal_count_diff": "출금 건수 차이",
    "txn_withdrawal_amount": "출금 금액(거래)",
    "sum_withdrawal_amount": "출금 금액(집계표)",
    "withdrawal_amount_diff": "출금 금액 차이",
    "opening_balance": "기초잔액",
    "closing_balance": "기말잔액(집계표)",
    "expected_closing_balance": "기말잔액(등식 계산)",
    "balance_diff": "잔액 차이",
    "deposit_mismatch": "입금 불일치",
    "withdrawal_mismatch": "출금 불일치",
    "balance_mismatch": "잔액 불일치",
}

STAT_LABELS = (
    ("total_rows", "총 거래 건수"),
    ("clean_rows", "정상 건수"),
    ("quarantined_rows", "격리 건수"),
    ("review_rows", "검토 대상 건수 (WARN, 정상에 포함)"),
    ("error_issues", "ERROR 오류 수"),
    ("warn_issues", "WARN 오류 수"),
    ("accounts_matched", "대사 일치 계좌 수"),
    ("accounts_mismatched", "대사 불일치 계좌 수"),
    ("accounts_excluded", "집계표 오류로 대사 제외 계좌 수"),
    ("accounts_existence_mismatch", "계좌 존재 불일치 수"),
    ("unidentified_quarantined_rows", "계좌 미식별 격리 건수"),
)


@dataclass(frozen=True)
class OutputPaths:
    report: Path
    clean_csv: Path
    quarantine_csv: Path


def write_outputs(
    out_dir: Path,
    run_id: int,
    base_date: date,
    loaded: Mapping[str, LoadedFile],
    result: ValidationResult,
    warnings: Sequence[str] = (),
) -> OutputPaths:
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = OutputPaths(
        report=out_dir / f"report_{run_id}.xlsx",
        clean_csv=out_dir / "clean_transactions.csv",
        quarantine_csv=out_dir / "quarantine_transactions.csv",
    )
    # 엑셀에서 바로 열 수 있도록 BOM이 있는 UTF-8로 저장한다.
    result.clean.to_csv(paths.clean_csv, index=False, encoding="utf-8-sig")
    result.quarantine.to_csv(paths.quarantine_csv, index=False, encoding="utf-8-sig")

    workbook = Workbook()
    _summary_sheet(workbook.active, run_id, base_date, loaded, result, warnings)
    file_names = {ft: loaded[ft].file_name for ft in FILE_TYPES}
    _issue_sheet(workbook.create_sheet("오류목록"), result.issues, file_names)
    _recon_sheet(workbook.create_sheet("대사결과"), result.recon)
    _issue_sheet(workbook.create_sheet("검토대상"), result.review, file_names)
    workbook.save(paths.report)
    return paths


def _set(cell, value) -> None:
    if value is pd.NA or (isinstance(value, float) and pd.isna(value)):
        value = None
    elif hasattr(value, "item"):
        value = value.item()
    cell.value = value
    # 데이터에 '='로 시작하는 값이 있어도 수식으로 실행되지 않게 문자열로 고정한다.
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"


def _write_rows(sheet, start_row: int, rows: Iterable[Sequence], header: Sequence[str] | None = None) -> int:
    row_idx = start_row
    if header is not None:
        for col_idx, name in enumerate(header, start=1):
            cell = sheet.cell(row=row_idx, column=col_idx)
            _set(cell, name)
            cell.font = Font(bold=True)
        row_idx += 1
    for values in rows:
        for col_idx, value in enumerate(values, start=1):
            _set(sheet.cell(row=row_idx, column=col_idx), value)
        row_idx += 1
    return row_idx


def _widths(sheet, widths: Sequence[int]) -> None:
    for idx, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = width


def _summary_sheet(sheet, run_id, base_date, loaded, result, warnings) -> None:
    sheet.title = "요약"
    rows: list[tuple] = [("실행 번호", run_id), ("기준일", base_date.isoformat())]
    rows += [(label, result.stats[key]) for key, label in STAT_LABELS]
    rows.append(("경고", "\n".join(warnings) if warnings else "없음"))
    next_row = _write_rows(sheet, 1, rows, header=("항목", "값"))

    next_row = _write_rows(
        sheet,
        next_row + 1,
        [
            (FILE_LABELS[ft], loaded[ft].file_name, loaded[ft].encoding, loaded[ft].row_count, loaded[ft].sha256)
            for ft in FILE_TYPES
        ],
        header=("입력 파일", "파일명", "인코딩", "행 수", "SHA-256"),
    )
    _write_rows(
        sheet,
        next_row + 1,
        [(r.rule_id, r.name, r.severity, result.rule_counts[r.rule_id]) for r in all_rules()],
        header=("규칙", "이름", "심각도", "오류 건수"),
    )
    _widths(sheet, (34, 26, 12, 10, 66))


def _issue_sheet(sheet, frame: pd.DataFrame, file_names: Mapping[str, str]) -> None:
    rows = (
        (r.rule_id, r.severity, file_names.get(r.file_type, r.file_type), r.row_no, r.record_key, r.column_name, r.value, r.message)
        for r in frame[list(ISSUE_COLUMNS)].itertuples(index=False)
    )
    _write_rows(sheet, 1, rows, header=ISSUE_HEADERS)
    sheet.freeze_panes = "A2"
    _widths(sheet, (8, 8, 22, 11, 20, 26, 20, 70))


def _recon_sheet(sheet, frame: pd.DataFrame) -> None:
    columns = list(RECON_COLUMNS)
    _write_rows(
        sheet,
        1,
        frame[columns].itertuples(index=False, name=None),
        header=[RECON_HEADERS[c] for c in columns],
    )
    sheet.freeze_panes = "B2"
    _widths(sheet, (16, 22, 26, *([14] * (len(columns) - 3))))
